import openpyxl
from datetime import datetime

print("=" * 60)
print("ANALISIS ESTRUCTURAL DE PLANILLAS — Cuarto Ambulante")
print(f"Ejecutado: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}")
print("=" * 60)

files = {
    "PLANTILLA.xlsx":        "Planilla de ventas diarias por pyme",
    "TOTAL_ABRIL_2026.xlsx": "Registro mensual consolidado",
    "PUNTO_DE_ENTREGA.xlsx": "Control de paquetes y retiros",
}

for fname, desc in files.items():
    try:
        wb = openpyxl.load_workbook(fname, data_only=True)
        print(f"\n[ARCHIVO] {fname}")
        print(f"  Descripcion : {desc}")
        print(f"  N° de hojas : {len(wb.sheetnames)}")
        print(f"  Hojas       : {', '.join(wb.sheetnames)}")
        for sn in wb.sheetnames:
            ws = wb[sn]
            nrows = sum(1 for r in ws.iter_rows(min_row=2, values_only=True) if any(v for v in r))
            print(f"    └─ [{sn}] → {nrows} filas con datos, {ws.max_column} columnas")
    except FileNotFoundError:
        print(f"\n[ERROR] No se encontro el archivo: {fname}")
        print("  Verifica que el script este en la misma carpeta que los .xlsx")

print("\n" + "=" * 60)
print("Script 1 finalizado.")
print("=" * 60)